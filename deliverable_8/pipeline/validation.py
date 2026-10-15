"""Fail on changed input data, inflated joins, arithmetic errors or lost IDs."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .sources import digest, input_sources


def coverage_table(joined, bonds, config):
    rows=[]
    for period in pd.period_range(config['start_month'],config['end_month'],freq='M'):
        month=str(period)
        frame=joined.loc[joined.snapshot_month.eq(month)]
        quarter=period.asfreq('Q').start_time
        rows.append({'snapshot_month':month,'quarter_start':quarter,
                     'snapshot_rows':len(frame), 'observed_prices':int(frame.airbnb_nightly_nzd.notna().sum()),
                     'missing_prices':int(frame.airbnb_nightly_nzd.isna().sum()),
                     'all_prices_missing':bool(len(frame) and frame.airbnb_nightly_nzd.isna().all()),
                     'assigned_location_rows':int(frame.area_code.notna().sum()),
                     'unassigned_location_rows':int(frame.area_code.isna().sum()),
                     'matched_bond_rows':int(frame.bond_row_matched.sum()),
                     'usable_price_gaps':int(frame.nightly_price_gap_nzd.notna().sum()),
                     'near_boundary_rows':int(frame.near_sa2_boundary.sum()),
                     'bond_quarter_available':bool(bonds.quarter_start.eq(quarter).any())})
    return pd.DataFrame(rows)


def validate_analysis(root, manifest, joined, bond_main, area_monthly, city_monthly, audit, areas, config):
    checks={}
    checks['source_hashes_unchanged']=all(digest(root / s['path'])==s['sha256'] for s in input_sources(manifest))
    checks['unique_city_month_segment_keys']=not city_monthly.duplicated(['snapshot_month','segment']).any()
    checks['city_monthly_nightly_arithmetic']=np.allclose(city_monthly.city_bond_nightly_equivalent_nzd,city_monthly.city_bond_weekly_median_nzd/7,equal_nan=True)
    checks['city_difference_of_medians_arithmetic']=np.allclose(city_monthly.citywide_difference_of_medians_nzd,city_monthly.airbnb_median_nightly_nzd-city_monthly.city_bond_nightly_equivalent_nzd,equal_nan=True)
    checks['no_city_benchmark_in_listing_or_area_records']=not any(c.startswith('city_bond_') for c in list(joined.columns)+list(area_monthly.columns))
    checks['unique_listing_month_keys']=not joined.duplicated(['listing_id','snapshot_month']).any()
    checks['only_christchurch_publisher_city']=joined.neighbourhood_group.str.strip().eq(config['city']).all()
    checks['only_new_zealand_airbnb_sources']=all(s['url'].startswith('https://data.insideairbnb.com/new-zealand/') for s in manifest['airbnb'])
    checks['unique_bond_area_quarter_keys']=not bond_main.duplicated(['area_code','quarter_start']).any()
    checks['no_listing_loss_or_inflation']=len(joined)==int(audit.retained_rows.sum())
    checks['correct_calendar_quarters']=joined.quarter_start.eq(joined.snapshot_date.dt.to_period('Q').dt.start_time).all()
    checks['valid_positive_prices']=joined.airbnb_nightly_nzd.dropna().gt(0).all()
    checks['valid_nightly_arithmetic']=np.allclose(joined.bond_nightly_equivalent_nzd,joined.bond_weekly_median_nzd/7,equal_nan=True)
    checks['valid_gap_arithmetic']=np.allclose(joined.nightly_price_gap_nzd,joined.airbnb_nightly_nzd-joined.bond_nightly_equivalent_nzd,equal_nan=True)
    checks['no_carry_forward_when_bond_missing']=joined.loc[~joined.bond_row_matched,'bond_nightly_equivalent_nzd'].isna().all()
    checks['no_false_gap_without_price']=joined.loc[joined.airbnb_nightly_nzd.isna(),'nightly_price_gap_nzd'].isna().all()
    checks['only_city_area_codes']=set(joined.area_code.dropna()).issubset(set(areas.area_code))
    checks['all_bond_rows_are_all_all']=bond_main.dwelling_type.eq('ALL').all() and bond_main.number_of_beds.eq('ALL').all()
    checks['unique_area_month_segment_keys']=not area_monthly.duplicated(['area_code','snapshot_month','segment']).any()
    checks['ranking_uses_observed_gap_count']=area_monthly.eligible_gap_ranking.eq(area_monthly.usable_gap_count.ge(config['minimum_priced_listings'])).all()
    reconciled=True
    for month in joined.snapshot_month.unique():
        original=joined.loc[joined.snapshot_month.eq(month)]
        summary=area_monthly.loc[area_monthly.snapshot_month.eq(month)&area_monthly.segment.eq('all_rooms')]
        reconciled &= int(summary.airbnb_listing_count.sum())==int(original.area_code.notna().sum())
    checks['area_counts_reconcile_with_assigned_listings']=reconciled
    checks['source_id_precision_preserved']=joined.listing_id.map(lambda x:isinstance(x,str) and x.isdigit()).all()
    checks={k:bool(v) for k,v in checks.items()}
    if not all(checks.values()):
        raise ValueError('Analysis checks failed: '+', '.join(k for k,v in checks.items() if not v))
    return checks


def verify_workbooks(paths, joined, bond_detail, bond_main, area_monthly, coverage, city_bonds, city_monthly):
    checks={}
    datasets=load_workbook(paths[1],read_only=True,data_only=True)
    expected={'Airbnb merged':len(joined),'Bond area quarter':len(bond_main),'Bond category detail':len(bond_detail)}
    for name,count in expected.items():
        checks[f'excel_{name}_row_count']=datasets[name].max_row==count+1
    ws=datasets['Airbnb merged']
    # Check every identifier after export, not just a numeric display sample.
    identifiers=[row[0] for row in ws.iter_rows(min_row=2,max_col=1,values_only=True)]
    checks['excel_all_listing_ids_exact_text']=identifiers==joined.listing_id.tolist() and all(isinstance(x,str) for x in identifiers)
    checks['excel_dataset_city_bond_rows']=datasets['City monthly bonds'].max_row==len(city_bonds)+1
    datasets.close()
    analysis=load_workbook(paths[0],read_only=True,data_only=True)
    checks['excel_area_monthly_row_count']=analysis['Area monthly'].max_row==len(area_monthly)+1
    checks['excel_coverage_row_count']=analysis['Coverage'].max_row==len(coverage)+1
    checks['excel_city_bond_rows']=analysis['City monthly bonds'].max_row==len(city_bonds)+1
    rows=list(analysis['Monthly trends'].values)
    fields=[f.replace(' ','_') for f in rows[0]]
    exported=pd.DataFrame(rows[1:],columns=fields)
    expected=city_monthly.loc[city_monthly.segment.eq('entire_short_stay')]
    checks['excel_citywide_gaps_exact']=np.allclose(pd.to_numeric(exported.citywide_difference_of_medians_nzd),expected.citywide_difference_of_medians_nzd,equal_nan=True)
    if 'Price recovery check' in analysis.sheetnames:
        check_rows=list(analysis['Price recovery check'].iter_rows(min_row=2,max_row=len(coverage)+1,values_only=True))
        checks['excel_price_source_audit_months']=[r[0] for r in check_rows]==coverage.snapshot_month.tolist()
        checks['excel_price_source_audit_values']=all(r[5]==0 for r in check_rows)
    analysis.close()
    if not all(checks.values()):
        raise ValueError('Exported workbook verification failed')
    return checks
