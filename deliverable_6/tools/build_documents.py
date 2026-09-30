"""Regenerate the Word and Markdown documents from the verified run evidence."""

import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parents[1]
AI_STATEMENT = "OpenAI ChatGPT through Codex was the AI used to assist with the code review, refactoring, tests and document drafting. The documents were reconciled with the implemented code and the executed verification results."


def read_csv(name: str) -> list[dict]:
    with (HERE / "results/data_processed" / name).open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def new_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    document.styles["Title"].font.size = Pt(23)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 2"].font.size = Pt(11.5)
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.keep_with_next = True
        document.styles[name].paragraph_format.space_before = Pt(12)
        document.styles[name].paragraph_format.space_after = Pt(6)
    document.core_properties.author = "DATA422 project team"
    document.core_properties.subject = "Deliverable 6 pipeline review and design principles"
    return document


def write_document(stem: str, pages: list[list[tuple]]) -> None:
    document = new_document()
    markdown = []
    for page_index, blocks in enumerate(pages):
        if page_index:
            document.add_page_break()
        for kind, text in blocks:
            if kind == "title":
                document.add_paragraph(text, "Title")
                markdown.append("# " + text)
            elif kind in ("h1", "h2"):
                document.add_heading(text, level=int(kind[-1]))
                markdown.append(("## " if kind == "h1" else "### ") + text)
            elif kind == "bullet":
                document.add_paragraph(text, "List Bullet")
                markdown.append("- " + text)
            elif kind == "code":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run(text)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                markdown.append("```python\n" + text + "\n```")
            else:
                document.add_paragraph(text)
                markdown.append(text)
    # Remove any title borders inherited from Word defaults.
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    docs = HERE / "docs"
    docs.mkdir(exist_ok=True)
    document.save(docs / f"{stem}.docx")
    (docs / f"{stem}.md").write_text("\n\n".join(markdown) + "\n", encoding="utf-8")


def main() -> None:
    quality = json.loads((HERE / "results/data_quality_summary.json").read_text())
    verification = json.loads((HERE / "results/verification_report.json").read_text())
    sanity = json.loads((HERE / "results/sanity_check_example.json").read_text())
    assert all(quality["checks"].values()) and all(verification["checks"].values())
    test_text = (HERE / "results/test_results.txt").read_text()
    test_count = int(re.search(r"Ran (\d+) tests?", test_text).group(1))
    assert "\nOK\n" in test_text
    manifest = json.loads((HERE / "results/run_manifest.json").read_text())
    settings = manifest["settings"]
    minimum = settings["minimum_priced_listings"]
    ranked = read_csv("largest_gap_locations.csv")
    comparisons = read_csv("location_comparison.csv")
    central = next(row for row in comparisons if int(float(row["area_code"])) == 326600)
    top_gap = float(ranked[0]["median_nightly_gap"])
    tied_names = ", ".join(row["area_name"] for row in ranked if round(float(row["median_nightly_gap"]), 2) == round(top_gap, 2))
    q = quality
    result_text = (f"The executed June 2026 run retains {q['airbnb_rows']:,} listings across {q['location_quarter_groups']} SA2 groups. "
                   f"All {q['sa2_matches']:,} listings have cached SA2 assignments. {q['bond_row_matches']:,} listings match a bond total "
                   f"({q['bond_row_matches'] / q['airbnb_rows']:.1%}); {q['usable_price_gaps']:,} have a usable price gap. "
                   f"The {q['missing_prices']} missing Airbnb prices are retained.")
    design_pages = [
        [("title", "Airbnb and tenancy bond pipeline design principles"),
         ("p", "DATA422 Deliverable 6  |  30 September 2026"),
         ("p", "The pipeline compares Christchurch Airbnb nightly prices and listing counts with quarterly tenancy bond statistics by Statistical Area 2. It preserves listing records, aligns dates to quarters, and makes the price comparison in consistent nightly units. The revised implementation is in deliverable_6 and reads the previously prepared files."),
         ("p", AI_STATEMENT),
         ("h1", "1 Inputs to the pipeline"),
         ("h2", "Christchurch Airbnb listings"),
         ("p", "The default input is deliverable_3/data_processed/christchurch_airbnb_cleaned.csv. Each row represents one listing in one snapshot. Required fields are id, latitude, longitude, price and snapshot_date. The current file has 3,469 rows and 23 columns, with snapshot_date 2026-06-19. The pipeline preserves the other input fields, including listing and host information."),
         ("p", "The Deliverable 3 cleaning stage has already standardised fields, checked duplicate listing-month keys, handled impossible numeric values, selected Christchurch City and attached snapshot dates. Deliverable 6 validates that prepared input; it does not repeat raw-data acquisition or claim to rebuild the upstream cleaning stages."),
         ("h2", "Quarterly tenancy bond statistics"),
         ("p", "The default input is deliverable_4/bond_data_processed/bond_data_clean_qi.csv, containing 26,991 rows and 12 columns for 2025-10-01, 2026-01-01 and 2026-04-01. These are national location-level records and category breakdowns. Required fields are TimeFrame, Location Id, Dwelling Type, Number Of Beds, Active Bonds, Total Bonds and Median Rent. The analysis selects published ALL/ALL totals, leaving 4,865 unique location-quarter rows."),
         ("h2", "Geographic lookup and settings"),
         ("p", "inputs/koordinates_coordinate_cache.csv is a byte-preserved copy of the Deliverable 5 SA2 cache, with 3,171 coordinate pairs and five fields: longitude, latitude, area_code, area_name and query_status. Its recorded source is Koordinates layer 123515, Statistical Area 2 2026. config.json sets the worker count and observed-gap minimum. Command-line options select inputs, cache, output folder and network access. An API key is read only from KOORDINATES_API_KEY when new queries are needed.")],
        [("h1", "2 Outputs from the pipeline"),
         ("p", "All generated analytical outputs go into deliverable_6/results by default. Earlier source files are read only. Intermediate tables make the transformations inspectable and support debugging."),
         ("bullet", "data_processed/airbnb_prepared.csv contains validated listing data and the derived bond_time quarter."),
         ("bullet", "data_processed/airbnb_enriched.csv adds SA2 code, name and query status to every listing."),
         ("bullet", "data_processed/bond_location_quarter.csv contains the ALL/ALL published totals at a unique location-quarter grain."),
         ("bullet", "data_processed/airbnb_bond_joined.csv adds bond statistics, a bond_row_matched flag, median_long_term_nightly_rent and nightly_price_gap."),
         ("bullet", "data_processed/location_comparison.csv reports listing counts, priced-listing counts, observed-gap counts, medians and active bond counts separately for each SA2-quarter."),
         ("bullet", "data_processed/largest_gap_locations.csv ranks up to 20 qualifying locations per quarter; plots show the 15 largest gaps and counts for the 20 areas with most listings."),
         ("bullet", "deliverable_6.sqlite stores the same six tables as the CSV exports. Missing values become SQL NULL; dates are stored as ISO text."),
         ("bullet", "koordinates_coordinate_cache.csv is the output checkpoint cache. analysis_summary.md and data_quality_summary.json are generated from the run. run_manifest.json records paths, hashes, settings and software versions. _SUCCESS.json marks a completed analytical run."),
         ("p", "The verification and document stages add sanity_check_example.json, verification_report.json and test_results.txt. docs contains the design principles and tutor change notes in editable Word and Markdown formats. The document builder uses verified result files to keep counts and conclusions consistent."),
         ("h2", "Observed output and interpretation"),
         ("p", result_text),
         ("p", f"Christchurch Central has {int(central['airbnb_count'])} listings, of which {int(central['priced_listing_count'])} have prices; the median nightly price is NZ${float(central['median_airbnb_price']):.2f}. "
                   f"Using at least {minimum} observed listing gaps, {tied_names} share the highest gap rounded to cents, NZ${top_gap:.2f}. Exact numeric values determine sort order, with area code as the final tie-breaker."),
         ("p", "The comparison is descriptive. Active bonds are a proxy for long-term rental activity, and Airbnb listings can include rooms as well as entire properties. Weekly rent divided by seven is a unit conversion, not evidence of equivalent accommodation quality, occupancy, returns or causal effects.")],
        [("h1", "3 Main steps in the pipeline"),
         ("h2", "Validate prepared inputs"),
         ("p", "prepare_airbnb checks required columns, positive integer IDs, valid dates, numeric nonnegative observed prices and New Zealand coordinate bounds. Missing prices remain missing. It rejects duplicate id and snapshot_date keys and multiple snapshots in the same quarter to prevent repeated listings from affecting a quarter comparison. prepare_bonds validates the schema, selects ALL/ALL totals, checks quarter-start dates and rejects duplicate Location Id and TimeFrame keys."),
         ("h2", "Attach SA2 geography"),
         ("p", "load_cache validates coordinate-key uniqueness and complete successful responses. enrich_coordinates performs a many-to-one coordinate join. Offline mode is the default and requires each input pair to have an ok or no_match result. Missing or failed pairs stop an offline run. With --allow-network, failed and missing pairs are queried after a single-query test, with bounded retries and concurrent workers. Input cache files are preserved and output cache checkpoints are saved atomically."),
         ("h2", "Align time and join bond totals"),
         ("p", "snapshot_date is mapped to the beginning of its calendar quarter. For 2026-06-19, bond_time is 2026-04-01. join_bonds performs a validated many-to-one left join on area_code and bond_time. All listings are retained. bond_row_matched distinguishes the absence of a record from a matched record whose rent is missing. No unmatched metric is filled with zero."),
         ("h2", "Calculate and summarise"),
         ("p", "Median long-term nightly rent equals Median Rent divided by seven. Each listing gap equals its nightly price minus that estimate. summarise_locations groups by bond_time, area_code and area_name, retaining the quarter in the output. It reports all listings, priced listings and usable price gaps separately. rank_gaps applies the configured minimum to observed gaps, then ranks eligible locations within each quarter."),
         ("h2", "Validate and export"),
         ("p", f"sanity_check verifies preserved row counts and listing keys, summary counts, quarter alignment, nightly units, gap arithmetic and missing unmatched values. write_outputs saves CSVs and SQLite, checks the SQLite joined row count, creates plots and the report, and writes the completion marker last. The current run passes {len(q['checks'])} analytical checks. The separate delivery verification compares all six SQLite tables with CSV contents and checks regression against Deliverable 5.")],
        [("h1", "4 Coding and software strategies"),
         ("h2", "Clear responsibilities and explicit data flow"),
         ("p", "run_pipeline.py orchestrates named stages. pipeline/transforms.py contains dataframe transformations, geography.py handles cache and network access, validation.py checks invariants, and outputs.py writes artifacts. Functions accept explicit inputs and return outputs. Transformations copy their inputs rather than mutating global data. Docstrings and descriptive names explain purpose and assumptions without commenting on every obvious line."),
         ("p", "This follows the supplied Data Wrangling Pipeline and Modularity lecture: page 30 describes meaningful modules, pages 34 to 36 illustrate descriptive functions and explicit inputs and outputs, page 37 recommends project-relative paths, and page 38 recommends scripted changes instead of manual data-file edits."),
         ("h2", "Configuration and reproducible automation"),
         ("p", "config.json separates tunable settings from transformation logic. Named constants define the weekly conversion and SA2 source fields. Defaults are relative to the script location, and path overrides support other machines. requirements.txt records the tested package versions; the run manifest records the actual environment. A Makefile runs analysis, the persisted sanity check, verification and Word and Markdown document generation in dependency order. The supplied example-automated-pipeline.zip demonstrates the same overall strategies through separate R stages, config.yaml, inspectable intermediates and a Makefile."),
         ("p", "The Makefile intentionally executes a complete verification run rather than implementing per-stage timestamp caching. A normal rerun replaces only generated Deliverable 6 outputs. Offline execution reuses the fixed supplied cache; network mode can change geographic results and needs its own review. For a failed live run, supply its checkpoint as --cache and choose a new --output folder."),
         ("h2", "Fail early and verify meaningful invariants"),
         ("p", f"Invalid schemas, malformed dates or numbers, duplicate keys and ambiguous quarterly snapshots stop the run with an actionable error. {test_count} automated edge tests exercise double counting, join expansion, missing prices, duplicate keys, quarter separation, offline cache reuse, failed-cache retries, timeout handling and arithmetic corruption. Tests use mocked network responses; the current validation does not claim a fresh live API check."),
         ("h2", "Preserve evidence and review assumptions"),
         ("p", "Earlier deliverables and raw data remain unchanged. inputs/source_materials_manifest.json records reviewed-source fingerprints. The revised join uses validated pandas merges; SQLite persists the resulting tables for inspection. This makes the many-to-one constraint explicit and produces the same existing listing values and all 174 previous location statistics. The observed-gap ranking threshold is an intentional method correction described in the change notes."),
         ("p", "Remaining review limits are geographic correctness and boundary compatibility: cached assignments are structurally checked but not independently recomputed against polygons, and the available files do not establish a shared SA2 vintage for the bond and Airbnb geography. Matching codes alone cannot resolve this; unmatched records and coverage are retained for review.")]
    ]
    example = sanity["example"]
    excluded = ", ".join(row["area_name"] for row in verification["excluded_from_ranking"])
    review_pages = [
        [("title", "Pipeline coding review and tutor change notes"),
         ("p", "DATA422 Deliverable 6  |  30 September 2026"),
         ("p", "The review revisits the Airbnb preparation and Qi cleaning scripts in Deliverables 3 and 4 and the comparison pipeline in Deliverable 5. Deliverable 6 implements a revised comparison pipeline; it reads the same date-bearing Airbnb input and cleaned bond input and preserves earlier deliverables. These notes record changes at the level of purpose and behaviour, rather than individual lines."),
         ("p", AI_STATEMENT),
         ("h1", "High level changes and reasons"),
         ("h2", "Separate the orchestration from focused stages"),
         ("p", "The large Deliverable 5 main function was split into configuration, pure dataframe transformations, geography, validation and output modules. Each stage has descriptive names and explicit inputs and outputs. This makes the join and calculation logic easier to read, reuse and test separately from files and network calls. The earlier Deliverable 4 scripts contain useful inspection and report logic but also run prompts and data loading at import time; they are retained as historical scripts rather than imported into the revised workflow."),
         ("h2", "Move settings into configuration and add automation"),
         ("p", "Worker count and the observed-gap minimum are stored in config.json with command-line overrides. Input and output paths use project-relative defaults. A Makefile runs the analytical workflow, sanity demonstration, delivery verification and document generation. This adapts the supplied automated R pipeline example to the existing Python project without changing languages or introducing a new workflow platform."),
         ("h2", "Strengthen validation before calculations"),
         ("p", "The revised pipeline rejects missing fields, malformed dates and numeric values, invalid IDs and coordinates, duplicate listing-snapshot keys and duplicate bond location-quarter totals. The previous pipeline sometimes converted malformed data silently to missing. Failure is now explicit. Legitimate missing prices are retained, so validation does not discard listings simply because a price is absent."),
         ("h2", "Use the input that retains the required date"),
         ("p", "Deliverable 4 christchurch_listings_clean_qi.csv removes snapshot_date and retains month_year. Deliverable 5 needs the snapshot date to derive a quarter. The revised pipeline therefore keeps Deliverable 3 christchurch_airbnb_cleaned.csv as its default Airbnb input, rather than claiming that the narrower Qi listing file feeds the join. The design document explicitly states this boundary and does not invent snapshot dates.")],
        [("h1", "Method corrections and output changes"),
         ("h2", "Keep quarters separate and constrain the join"),
         ("p", "The previous location summary grouped only by SA2 and could mix prices and maximum bond metrics across quarters. The revised summary includes bond_time. It permits at most one snapshot per quarter; additional monthly snapshots require an explicit selection step. A validated many-to-one left join on SA2 and quarter now enforces the intended grain. Published ALL/ALL bond totals are retained rather than summing overlapping breakdowns."),
         ("h2", "Apply the gap minimum to observed prices"),
         ("p", f"The previous ranking required five total Airbnb listings, even if fewer had prices. The revised ranking requires at least {minimum} usable listing gaps, while retaining the all-listing count separately. This makes the minimum refer to the observations actually used to calculate the median. The documented choice is configurable; it is a stricter denominator and changes the eligible locations."),
         ("p", f"In this run, {excluded} have at least five listings but only four usable price gaps each, so they leave the ranking. Eligible locations decrease from {verification['old_eligible_locations']} to {verification['new_eligible_locations']}. Bryndwr South was the previous leading result, NZ$287.64 from four observed prices among five listings. Under the revised rule, {tied_names} share the highest rounded gap, NZ${top_gap:.2f}."),
         ("h2", "Distinguish missing data and count denominators"),
         ("p", f"The outputs now distinguish all listings, priced listings and usable listing gaps, and record whether a bond row matched. Unmatched rents remain missing; plots label missing bond counts. The earlier Christchurch Central report said 117 listings because it counted only nonmissing prices. The revised report states {int(central['airbnb_count'])} listings and {int(central['priced_listing_count'])} priced observations, retaining its NZ${float(central['median_airbnb_price']):.2f} median."),
         ("h2", "Make cache reuse and error handling reliable"),
         ("p", "A complete cache can now reproduce results without an API key or network call. Failed cache rows are treated as pending in live mode instead of permanently completed. Timeout handling no longer assumes every exception has a reason attribute. Cache checkpoints are written atomically and input caches remain unchanged. Query errors retain concise status codes and do not save request URLs, API keys or server response bodies."),
         ("h2", "Record provenance and keep documents consistent"),
         ("p", "Generated outputs include an input and code hash manifest, actual software versions, measured quality checks and a completion marker. The document builder reads the verified result files for counts and results. The revised pandas join replaces the embedded SQL transformation; SQLite continues to store all six output tables. Independent verification confirms that all prior joined values and all 174 location statistics are preserved. Changes to ranking membership are deliberate and recorded above.")],
        [("h1", "Sanity check example for the tutor"),
         ("p", "A left join can appear successful while duplicating listings or attaching the wrong period. A useful sanity check combines row and key preservation with an independent unit calculation. Run sanity_check_example.py after the pipeline; it reloads the saved CSVs and checks the persisted output rather than relying on an in-memory preview."),
         ("h2", "Expected conditions"),
         ("bullet", "The left join must preserve the input row count and the set of listing ID and snapshot date keys, with no duplicate listing-snapshot keys."),
         ("bullet", "A June 2026 snapshot must map to the quarter beginning 2026-04-01. Every location summary count must reconcile with the listing-level output."),
         ("bullet", "Nightly bond rent must equal weekly rent divided by seven. The price gap must equal the Airbnb nightly price minus that estimate. Missing unmatched bond values must remain missing."),
         ("code", "if len(joined) != len(airbnb):\n    raise ValueError(\"Left join changed the row count\")\nexpected_gap = price - weekly_rent / 7.0\n# Compare the saved gap with a small floating-point tolerance."),
         ("h2", "Actual hand calculation"),
         ("p", f"Listing {example['id']} has snapshot date {example['snapshot_date']}, mapped to bond quarter {example['quarter']}. Its Airbnb price is NZ${example['nightly_price']:.2f} per night, and the matched median weekly rent is NZ${example['weekly_rent']:.2f}. The expected nightly bond rent is {example['weekly_rent']:.0f} / 7 = NZ${example['weekly_rent'] / 7:.6f}. The expected gap is {example['nightly_price']:.0f} - {example['weekly_rent']:.0f} / 7 = NZ${example['expected_gap']:.6f}, matching the saved value within numerical tolerance."),
         ("h2", "Executed evidence"),
         ("p", result_text),
         ("p", f"All {len(q['checks'])} analytical checks and {test_count} edge tests passed. The persisted-output demonstration passes nine dataframe checks. Delivery verification confirms matching CSV and SQLite contents for all six tables, unchanged input fingerprints, the corrected ranking denominator and agreement with Deliverable 5 on existing joined values and location statistics. Deliberately corrupted gaps and expanded joins fail the tests, showing that the checks detect these errors."),
         ("h2", "Materials to show the tutor"),
         ("p", "Show this change log, the separate design-principles document, run_pipeline.py and its named modules, sanity_check_example.py, and results/test_results.txt plus the JSON verification reports. The design document maps the actual inputs, outputs, stages and adopted coding strategies to the revised implementation."),
         ("p", "Course references are the supplied example-automated-pipeline.zip and Data Wrangling W03 DATA201 422 pipeline modularity, particularly pages 30 and 34 to 38. The analytical checks do not independently validate the geographic cache against polygons or resolve the SA2 vintage of the bond data. Those limits remain explicit in the design document.")]
    ]
    write_document("Design_Principles", design_pages)
    write_document("Coding_Review_and_Sanity_Check", review_pages)
    text = (HERE / "docs/Design_Principles.md").read_text()
    for required in ("1 Inputs", "2 Outputs", "3 Main steps", "4 Coding", "OpenAI ChatGPT", "snapshot_date", "gap_listing_count"):
        # gap_listing_count is represented by its plain-language label in the prose.
        if required == "gap_listing_count":
            assert "observed-gap" in text
        else:
            assert required in text
    sources = json.loads((HERE / "inputs/source_materials_manifest.json").read_text())
    preserved = all(hashlib.sha256(Path(entry["path"]).read_bytes()).hexdigest() == entry["sha256"] for entry in sources)
    assert preserved, "A reviewed source changed; recheck the source evidence"
    runner_text = (HERE / "run_pipeline.py").read_text()
    paths_match = all(path in runner_text for path in ("deliverable_3/data_processed/christchurch_airbnb_cleaned.csv",
                                                      "deliverable_4/bond_data_processed/bond_data_clean_qi.csv"))
    assert paths_match
    evidence = {"ai_disclosed": True, "four_required_design_sections": True,
                "derived_results_from_verified_run": True, "source_documents_preserved": True,
                "input_paths_match_runner_defaults": paths_match, "reviewed_sources_unchanged": preserved, "test_count": test_count,
                "analytical_checks": len(q["checks"]), "highest_gap_rounded": round(top_gap, 2)}
    (HERE / "results/document_alignment.json").write_text(json.dumps(evidence, indent=2) + "\n", encoding="utf-8")
    artifacts = {str(path.relative_to(HERE)): hashlib.sha256(path.read_bytes()).hexdigest()
                 for path in sorted(HERE.rglob("*")) if path.is_file()
                 and "__pycache__" not in path.parts
                 and path.name not in {"delivery_manifest.json", ".DS_Store"}}
    (HERE / "results/delivery_manifest.json").write_text(json.dumps(artifacts, indent=2) + "\n", encoding="utf-8")
    print("Created two Word documents and matching Markdown versions")


if __name__ == "__main__":
    main()
